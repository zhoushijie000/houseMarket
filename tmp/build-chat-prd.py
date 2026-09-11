from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

root=Path(__file__).resolve().parent.parent
d=Document(root/'成都住建房产超市直播专区需求.docx')
body=d._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'): body.remove(child)
sec=d.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    if name not in d.styles:d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s=d.styles[name];s.font.name='Microsoft YaHei';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt({'Normal':10,'Title':20,'Heading 1':15,'Heading 2':12,'Heading 3':11}[name])
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.line_spacing=1.15
def p(text):return d.add_paragraph(text,'Normal')
def h(text,level=2):return d.add_heading(text,level)
def items(lines):
    for i,t in enumerate(lines,1):p(f'{i}）{t}')
def fields(rows):
    t=d.add_table(rows=1,cols=3);t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for c in t.columns:c.width=Cm(17/3)
    for c,text in zip(t.rows[0].cells,['字段','定义','备注']):
        c.text=text
        for r in c.paragraphs[0].runs:r.bold=True
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');c._tc.get_or_add_tcPr().append(shade)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=text
    for row in t.rows:
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
                for r in para.runs:r.font.size=Pt(9)
    p('')

d.add_paragraph('成都住建房产超市在线沟通需求','Title')
p('版本：V1.0    编写日期：2026年9月10日')
p('系统：成都住建房产超市    产品：/    评审时间：待定')
h('1 功能清单',1)
p('小程序端：消息列表、在线聊天、置业顾问详情。用户可直接与顾问沟通、交换微信、电话联系，并自主屏蔽顾问消息。')
h('2 需求详情',1)
h('2.1 消息列表')
p('入口：底部导航“消息”。页面标题为“消息通知”。')
h('2.1.1 页面交互',3)
items(['列表展示顾问会话及平台通知；顾问姓名与“置业顾问”标签分开显示，楼盘名称显示在下方。','点击顾问会话进入在线聊天，无需先获取手机号；点击平台通知查看通知详情。','显示最近消息摘要和未读数量；头像缺失时使用姓名首字占位。时间取最近一条消息发送时间：今天显示 HH:mm，昨天显示“昨天 HH:mm”，昨天以前显示 YYYY-MM-DD；按北京时间自然日判断。'])
h('2.1.2 字段说明',3)
fields([
('头像','会话对象头像','图片；缺失时显示姓名首字。'),
('姓名与身份标签','顾问姓名及身份','例：李晓雨；“置业顾问”为独立标签。'),
('楼盘名称','顾问对应的楼盘','姓名下方展示；平台通知不展示。'),
('最近消息','最近一条消息的摘要','摘要超长省略；图片消息显示“[图片]”。'),
('消息时间','最近一条消息的发送时间','今天 HH:mm；昨天“昨天 HH:mm”；更早 YYYY-MM-DD。按北京时间自然日判断，收发新消息后更新。'),
('未读数量','当前会话未读消息条数','无未读时不显示；清零时机待确认。')])

d.add_page_break();h('2.2 在线聊天',1)
p('入口：消息列表中的顾问会话，或楼盘详情的在线联系入口。')
h('2.2.1 页面交互',3)
items(['顶部沿用统一导航，标题为“消息通知”；下方展示头像、姓名、身份标签和楼盘。点击头像或姓名进入顾问详情，返回后保留聊天内容。','用户消息靠右；顾问每条消息左侧展示头像，包括文字、二维码和微信号。输入非空文字后点击“发送”；空白内容不发送。','输入栏依次为输入框、发送、“＋”。点击“＋”在下方展开图片、拍照入口；再次点击或点击外部收起。选择图片或完成拍摄后发入会话，取消选择不发送。','点击“交换微信”，顾问依次发送微信二维码图片、微信号消息；微信号附“复制微信号”按钮，成功提示“已复制”，可在其他应用粘贴。复制失败提示“请长按微信号复制”。','点击“电话联系”调起该顾问号码的拨号界面；无号码显示“暂无电话”，不可拨打。','不展示“咨询对应房源”快捷按钮。屏蔽顾问后不再接收其新回复，已有消息保留。'])
h('2.2.2 字段说明',3)
fields([
('顾问信息','头像、姓名、身份标签、对应楼盘','与列表及顾问详情一致；楼盘不加“所属项目”前缀。'),
('消息内容与时间','双方文字或图片及发送时间','区分发送方；顾问消息逐条展示头像。'),
('输入内容','用户待发送文字','去除首尾空格；空内容不可发送。'),
('图片或照片','选择相册图片或调用相机拍摄','支持图片文件；设备需支持相机调用。'),
('微信二维码','当前顾问的微信二维码图片','需使用真实二维码；缺失时不得使用其他顾问二维码。'),
('微信号','当前顾问微信账号','支持复制；无数据时显示“暂未提供”。'),
('联系电话','当前顾问可联系号码','用于调起拨号；不直接自动拨出。')])

d.add_page_break();h('2.3 置业顾问详情',1)
p('入口：聊天顶部的顾问头像或姓名。页面标题为“置业顾问详情”。')
h('2.3.1 页面交互',3)
items(['展示顾问头像、姓名、“置业顾问”标签、楼盘名称和活跃情况，下方展示今日回复数、已服务客户数。','在详情页提供“屏蔽该置业顾问”开关，聊天页不再展示。默认关闭，并保留用户上次选择。','开启后提示“已屏蔽，您将不再接收该置业顾问的消息。”；停止接收该顾问新消息，其他顾问不受影响。','关闭后恢复接收，不显示“已取消屏蔽”提示。常驻说明为“开启后，您将不再接收该置业顾问的消息；关闭后恢复接收。”','返回聊天保留当前会话。统计卡不展示“原型示例数据”说明。'])
h('2.3.2 字段说明',3)
fields([
('顾问头像','顾问个人头像','与聊天一致；缺失时用姓名首字占位。'),
('姓名与身份标签','顾问姓名及“置业顾问”标签','姓名不拼接身份文字。'),
('楼盘名称','顾问服务的对应楼盘','显示在姓名下方。'),
('活跃情况','顾问近期使用情况及最近活跃时间','例：今日活跃、2分钟前；无数据展示“暂无数据”。'),
('今日回复数','当日顾问回复消息数量','整数；统计时区及计数口径待确认。'),
('已服务客户数','顾问累计服务客户数量','整数；客户去重和服务认定口径待确认。'),
('屏蔽状态','当前用户是否屏蔽该顾问','开关；按用户与顾问关系保存。')])
h('3 验收与待确认',2)
p('验收：可直接进入聊天；图文可发送；微信号可复制；电话可调起；顾问详情可进可退；屏蔽开关可保存并正确控制新消息接收。')
p('待确认：活跃和统计指标口径、未读清零规则、图片大小上限及发送失败重试规则。正式上线需接入真实消息、顾问资料和屏蔽状态服务；当前示例数据与二维码不作为真实业务数据。')

d.add_page_break();h('4 页面参考',1)
p('图1 消息列表    图2 在线聊天    图3 置业顾问详情')
line=d.add_paragraph()
for name in ['list','chat','advisor']:
    file=root/f'tmp/chat-prd-{name}.png'
    if file.exists():line.add_run().add_picture(str(file),width=Cm(5.5))
    else:p(f'{name} 页面截图待补充，参照当前在线沟通原型。')
media=root/'tmp/chat-prd-media.png'
if media.exists():
    p('图4 点击“＋”展开图片与拍照')
    d.add_picture(str(media),width=Cm(9))
d.core_properties.title='成都住建房产超市在线沟通需求'
d.core_properties.subject='消息列表 在线聊天 置业顾问详情'
out=root/'output/成都住建房产超市_在线沟通需求_V1.0.docx'
d.save(out)
assert all([c.text for c in t.rows[0].cells]==['字段','定义','备注'] for t in d.tables)
print(out)
