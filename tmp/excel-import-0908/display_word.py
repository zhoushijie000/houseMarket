from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out=Path(r'C:/Users/ZhuanZ/Desktop/zhou/houseMarket/outputs/import-template-0908/字段及相册小程序展示位置说明模板.docx')
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2)
sec.left_margin=sec.right_margin=Cm(2.3)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[name]; s.font.name='宋体'; s.font.size=Pt(11); s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
 s.paragraph_format.space_after=Pt(8)
 s.paragraph_format.line_spacing=1.25
doc.styles['Title'].font.size=Pt(18)
doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(14)
doc.styles['Heading 1'].font.bold=True
doc.add_paragraph('字段及相册小程序展示位置说明',style='Title')
doc.add_paragraph('按字段或相册逐项填写说明，并插入对应的小程序截图。新增内容时，复制相应模板页即可。')
def field(label):
 p=doc.add_paragraph();p.add_run(label+'：').bold=True;p.add_run('【填写】')
def page(kind):
 doc.add_paragraph(kind+'说明',style='Heading 1')
 for label in [kind+'名称','内容说明','小程序页面','具体展示位置']:
  field(label)
 p=doc.add_paragraph();p.add_run('展示位置截图').bold=True
 table=doc.add_table(rows=1,cols=1);table.alignment=WD_TABLE_ALIGNMENT.CENTER
 table.autofit=False;table.columns[0].width=Cm(16.4)
 row=table.rows[0];row.height=Cm(10);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
 cell=row.cells[0];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
 run=cell.paragraphs[0].add_run('【在此插入小程序截图】');run.font.color.rgb=RGBColor.from_string('888888')
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 table._tbl.tblPr.append(borders)
 doc.add_paragraph('')
 field('备注')
page('字段')
doc.add_page_break()
page('相册')
doc.save(out)
print(out)
