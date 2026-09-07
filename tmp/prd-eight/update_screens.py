from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from zipfile import ZipFile,ZIP_DEFLATED
from PIL import Image
from copy import deepcopy
from io import BytesIO

root=Path(__file__).resolve().parents[2]; tmp=Path(__file__).parent
src=root/'output/成都住建房产超市_小程序及管理端调整需求_V1.0_20260907.docx'
out=root/'output/成都住建房产超市_小程序及管理端调整需求_V1.1_含页面截图.docx'
d=Document(src)
groups=[
 ['01-楼盘导航','02-售楼部导航'],
 ['03-最新预售证','04-全部预售证','05-取证详情','06-一房一价','07-剩余房源'],
 ['08-报告摘要','10-楼盘报告'],
 ['11-我的','12-资料设置'],
 ['09-联系置业顾问','13-顾问对话'],
 ['14-地址维护','15-地图选点'],
 ['16-批量价格列表','17-价格导入','18-批量比例设置'],
 ['19-项目报告列表','20-项目报告录入']]
for p in list(d.paragraphs):
 if '验收要点' in p.text:
  n=p._p.getnext()
  if n is not None:p._p.getparent().remove(n)
  p._p.getparent().remove(p._p)
 elif p.text.startswith('联动验收：'):p.text=p.text.replace('联动验收：','跨端联动：',1)
 elif p.text.startswith('版本：'):p.text='版本：V1.1（补充页面截图）'

placeholders=[p for p in d.paragraphs if p.text.startswith('页面截图：待补充')]
assert len(placeholders)==8
for section,(p,names) in enumerate(zip(placeholders,groups),1):
 p.text='页面截图（当前原型，图中数据为演示数据）'
 p.paragraph_format.keep_with_next=True
 cursor=p._p
 for i,name in enumerate(names,1):
  path=tmp/'screens'/f'{name}.png'
  with Image.open(path) as im:w,h=im.size
  width=min(7.5 if section>=6 else 3.15,6.4*w/h)
  pic=d.add_paragraph();pic.alignment=WD_ALIGN_PARAGRAPH.CENTER
  pic.paragraph_format.keep_with_next=True
  pic.paragraph_format.space_after=Pt(3)
  pic.add_run().add_picture(str(path),width=Inches(width))
  cursor.addnext(pic._p);cursor=pic._p
  cap=d.add_paragraph(style='dingdocnormal');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
  cap.paragraph_format.keep_with_next=False;cap.paragraph_format.space_after=Pt(10)
  cap.add_run(f'图2.{section}-{i} {name.split("-",1)[1]}').font.size=Pt(10)
  cursor.addnext(cap._p);cursor=cap._p
buf=BytesIO();d.save(buf);buf.seek(0)
with ZipFile(src) as z:parts={n:z.read(n) for n in z.namelist()}
with ZipFile(buf) as z:
 for n in z.namelist():
  if n in ['word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'] or n.startswith('word/media/'):parts[n]=z.read(n)
with ZipFile(out,'w',ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
final=Document(out)
assert len(final.inline_shapes)==20
assert not any('验收要点' in p.text or '待补充对应页面截图' in p.text for p in final.paragraphs)
assert len(final.tables)==19
print(out)
print('20 screenshots; 8 acceptance sections removed; 19 tables retained')
